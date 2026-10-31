# -*- coding: utf-8 -*-
"""Extract full text of manuscript_BIB_protocol_v2.docx (paragraphs + tables) to a UTF-8 txt."""
import docx

SRC = r"C:\D\workbuddy\科研\共病分析\课题设计阶段\manuscript_BIB_protocol_v2.docx"
OUT = r"C:\Temp\iad_g1\bib_v2_extracted.txt"

d = docx.Document(SRC)
lines = []
for p in d.paragraphs:
    t = p.text.strip()
    if t:
        lines.append(t)
for i, tbl in enumerate(d.tables):
    lines.append(f"\n=== TABLE {i+1} ===")
    for row in tbl.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        lines.append(" | ".join(cells))
text = "\n".join(lines)
with open(OUT, "w", encoding="utf-8") as f:
    f.write(text)
print(f"OK: {len(text)} chars, {len(lines)} lines -> {OUT}")
