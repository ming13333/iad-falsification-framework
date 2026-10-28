#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build a submission-ready DOCX from the Frontiers Methods manuscript.

Design decisions
----------------
* Pure python-docx. No pandoc, no LibreOffice, no editor SDK.
* Three-line (booktabs-style) tables: no vertical rules, one rule above the
  header, one below the header, one below the last row.
* Inline LaTeX is rendered into native Word runs: single latin letters italic,
  operators and multi-letter macros upright, superscripts and subscripts set
  with real Word formatting. The output contains no $ or backslash commands.
* Double line spacing, continuous line numbers and a page-number footer, which
  is what Frontiers asks for in a submitted manuscript file.

Usage
-----
    python code/md_to_docx_frontiers.py \
        --in  <manuscript.md> \
        --out <manuscript.docx>
"""

import argparse
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
BODY_SIZE = 12
TABLE_SIZE_WIDE = 8
TABLE_SIZE = 9
NOTE_SIZE = 9

# Hand-tuned column widths (cm) for the two tables where the even allocation
# starves a short label column. Total is 15.9 cm, the text width of A4 with
# 2.54 cm margins.
FIXED_WIDTHS = {
    1: [1.05, 2.15, 3.45, 3.05, 3.15, 3.05],   # Table 1, six gates
    4: [0.95, 2.25, 8.60, 2.20, 1.90],          # Table 4, protocol steps
}

# --------------------------------------------------------------------------
# LaTeX -> unicode
# --------------------------------------------------------------------------

GREEK = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε",
    "varepsilon": "ε", "zeta": "ζ", "eta": "η", "theta": "θ", "iota": "ι",
    "kappa": "κ", "lambda": "λ", "mu": "μ", "nu": "ν", "xi": "ξ", "pi": "π",
    "rho": "ρ", "sigma": "σ", "tau": "τ", "upsilon": "υ", "phi": "φ",
    "varphi": "φ", "chi": "χ", "psi": "ψ", "omega": "ω",
    "Gamma": "Γ", "Delta": "Δ", "Theta": "Θ", "Lambda": "Λ", "Xi": "Ξ",
    "Pi": "Π", "Sigma": "Σ", "Upsilon": "Υ", "Phi": "Φ", "Psi": "Ψ",
    "Omega": "Ω",
}

SYMBOLS = {
    "approx": "≈", "sim": "~", "simeq": "≃", "propto": "∝", "equiv": "≡",
    "times": "×", "cdot": "·", "div": "÷", "pm": "±", "mp": "∓",
    "ge": "≥", "geq": "≥", "le": "≤", "leq": "≤", "ne": "≠", "neq": "≠",
    "ll": "≪", "gg": "≫",
    "rightarrow": "→", "to": "→", "leftarrow": "←", "leftrightarrow": "↔",
    "Rightarrow": "⇒", "Leftarrow": "⇐",
    "in": "∈", "notin": "∉", "subset": "⊂", "cup": "∪", "cap": "∩",
    "ldots": "…", "dots": "…", "cdots": "⋯",
    "infty": "∞", "partial": "∂", "sum": "Σ", "prod": "Π", "int": "∫",
    "sqrt": "√", "prime": "′", "ast": "∗", "star": "⋆", "circ": "∘",
    "lvert": "|", "rvert": "|", "vert": "|", "mid": "|", "|": "‖",
    " ": " ", ",": " ", ";": " ", ":": " ", "!": "",
    "quad": " ", "qquad": "  ",
}

UPRIGHT_MACROS = ("text", "mathrm", "mathbf", "operatorname", "textnormal")


def _read_group(s, i):
    """s[i] == '{' -> return (inner, index after the matching '}')."""
    depth = 0
    j = i
    while j < len(s):
        if s[j] == "\\":
            j += 2
            continue
        if s[j] == "{":
            depth += 1
        elif s[j] == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    return s[i + 1:], len(s)


def parse_math(expr, base="it"):
    """Return a list of (text, kind) with kind in {'rm','it','sup','sub'}.

    ``base`` is the style for a bare latin letter: italic at the top level,
    because a single letter in a formula is a variable, and upright inside
    \\text{...}, because that macro marks a multi-letter label.
    """
    out = []
    _emit(expr, base, out)
    # merge adjacent runs of the same kind first, then normalise dashes, so that
    # a dash pair split across two single-character tokens is still caught
    merged = []
    for text, kind in out:
        if merged and merged[-1][1] == kind:
            merged[-1] = (merged[-1][0] + text, kind)
        else:
            merged.append((text, kind))
    fixed = []
    for text, kind in merged:
        if kind == "rm":
            text = text.replace("---", "—").replace("--", "–")
        if text:
            fixed.append((text, kind))
    return fixed


def _emit(s, base, out):
    i = 0
    while i < len(s):
        c = s[i]
        if c == "\\":
            m = re.match(r"\\([A-Za-z]+)", s[i:])
            if m:
                cmd = m.group(1)
                i += 1 + len(cmd)
                if cmd in UPRIGHT_MACROS:
                    if i < len(s) and s[i] == "{":
                        grp, i = _read_group(s, i)
                        _emit(grp, "rm", out)
                    continue
                if cmd in ("left", "right"):
                    continue
                if cmd in GREEK:
                    out.append((GREEK[cmd], "it"))
                    continue
                if cmd in SYMBOLS:
                    out.append((SYMBOLS[cmd], "rm"))
                    continue
                # unknown macro: drop the backslash, keep the name upright
                out.append((cmd, "rm"))
                continue
            # escaped single character, e.g. \{ \} \, \%
            if i + 1 < len(s):
                nxt = s[i + 1]
                out.append((" " if nxt in ",;:" else nxt, "rm"))
                i += 2
                continue
            i += 1
            continue
        if c in "^_":
            i += 1
            if i < len(s) and s[i] == "{":
                grp, i = _read_group(s, i)
            elif i < len(s):
                grp = s[i]
                i += 1
            else:
                grp = ""
            inner = []
            _emit(grp, "it", inner)
            kind = "sup" if c == "^" else "sub"
            for t, _k in inner:
                out.append((t, kind))
            continue
        if c == "{":
            grp, i = _read_group(s, i)
            _emit(grp, base, out)
            continue
        if c == "}":
            i += 1
            continue
        if c.isascii() and c.isalpha():
            out.append((c, base))
            i += 1
            continue
        out.append((c, "rm"))
        i += 1


# --------------------------------------------------------------------------
# inline markdown -> runs
# --------------------------------------------------------------------------

def add_runs(paragraph, text, bold=False, italic=False, size=None,
             font=BODY_FONT, color=None):
    """Parse **bold**, *italic*, `code`, $math$ and emit Word runs."""
    i = 0
    buf = []

    def flush():
        if buf and "".join(buf):
            r = paragraph.add_run("".join(buf))
            r.bold = bold
            r.italic = italic
            r.font.name = font
            _east_asia(r, font)
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        buf.clear()

    while i < len(text):
        ch = text[i]
        if ch == "$":
            j = text.find("$", i + 1)
            if j > i:
                flush()
                for piece, kind in parse_math(text[i + 1:j]):
                    r = paragraph.add_run(piece)
                    r.bold = bold
                    r.italic = italic or kind == "it"
                    r.font.name = font
                    _east_asia(r, font)
                    if size:
                        r.font.size = Pt(size)
                    if color:
                        r.font.color.rgb = color
                    if kind == "sup":
                        r.font.superscript = True
                    elif kind == "sub":
                        r.font.subscript = True
                i = j + 1
                continue
        if text.startswith("**", i):
            flush()
            bold = not bold
            i += 2
            continue
        if ch == "*":
            flush()
            italic = not italic
            i += 1
            continue
        if ch == "`":
            j = text.find("`", i + 1)
            if j > i:
                flush()
                r = paragraph.add_run(text[i + 1:j])
                r.font.name = MONO_FONT
                _east_asia(r, MONO_FONT)
                r.font.size = Pt((size or BODY_SIZE) - 1)
                r.bold = bold
                r.italic = italic
                i = j + 1
                continue
        buf.append(ch)
        i += 1
    flush()


def _east_asia(run, font):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), font)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)


# --------------------------------------------------------------------------
# document scaffolding
# --------------------------------------------------------------------------

def setup_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.54)
    sec.right_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)

    add_line_numbers(sec)
    add_page_number(sec)


def add_line_numbers(section):
    sectPr = section._sectPr
    for el in sectPr.findall(qn("w:lnNumType")):
        sectPr.remove(el)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    cols = sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(ln)
    else:
        sectPr.append(ln)


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    p._p.append(fld)


def para(doc, text="", size=BODY_SIZE, bold=False, italic=False,
         align=None, space_before=0, space_after=6, line=WD_LINE_SPACING.DOUBLE,
         indent_cm=None, hanging_cm=None, keep_with_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if line is not None:
        pf.line_spacing_rule = line
    if align is not None:
        p.alignment = align
    if indent_cm is not None:
        pf.left_indent = Cm(indent_cm)
    if hanging_cm is not None:
        pf.first_line_indent = Cm(-hanging_cm)
    if keep_with_next:
        pf.keep_with_next = True
    if text:
        add_runs(p, text, bold=bold, italic=italic, size=size)
    return p


# --------------------------------------------------------------------------
# tables
# --------------------------------------------------------------------------

def _borders_none(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    for el in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(el)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")


def _cell_border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = OxmlElement("w:" + edge)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000")
    borders.append(el)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def parse_table(lines):
    rows = []
    for ln in lines:
        s = ln.strip()
        if not s.startswith("|"):
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue
        rows.append(cells)
    return rows


def column_widths(rows, total_cm=15.9, min_cm=1.5):
    ncols = max(len(r) for r in rows)
    weights = []
    for c in range(ncols):
        lens = []
        for r in rows:
            cell = r[c] if c < len(r) else ""
            plain = re.sub(r"\$[^$]*\$", "xxxxxxxx", cell)
            plain = re.sub(r"[*`]", "", plain)
            lens.append(len(plain))
        mean = sum(lens) / max(1, len(lens))
        weights.append(max(2.0, mean) ** 0.75)
    total = sum(weights)
    widths = [max(min_cm, total_cm * w / total) for w in weights]
    scale = total_cm / sum(widths)
    return [w * scale for w in widths]


def add_table(doc, rows, fixed_widths=None):
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    size = TABLE_SIZE_WIDE if ncols >= 5 else TABLE_SIZE
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _borders_none(table)

    widths = fixed_widths if fixed_widths else column_widths(rows)
    assert abs(len(widths) - ncols) < 1, "width vector does not match column count"

    tblPr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(round(sum(widths) * 567))))  # cm -> twips
    tblPr.insert_element_before(
        tblW, "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
        "w:tblLayout", "w:tblCellMar", "w:tblLook")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.insert_element_before(layout, "w:tblCellMar", "w:tblLook")

    for ri, rowdata in enumerate(rows):
        row = table.rows[ri]
        _cant_split(row)
        if ri == 0:
            _repeat_header(row)
        for ci, val in enumerate(rowdata):
            cell = row.cells[ci]
            cell.width = Cm(widths[ci])
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.space_before = Pt(2)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.0
            add_runs(p, val, bold=(ri == 0), size=size)
    # three-line rules
    for cell in table.rows[0].cells:
        _cell_border(cell, "top", 12)
        _cell_border(cell, "bottom", 8)
    for cell in table.rows[-1].cells:
        _cell_border(cell, "bottom", 12)
    # a thin spacer keeps the table off the following caption note
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(8)
    return table


# --------------------------------------------------------------------------
# main conversion
# --------------------------------------------------------------------------

def build(in_path, out_path):
    text = open(in_path, encoding="utf-8").read()
    lines = text.split("\n")
    doc = Document()
    setup_document(doc)

    # ---- front matter ----------------------------------------------------
    title = lines[0].lstrip("#").strip()
    p = para(doc, "", space_after=12)
    add_runs(p, title, bold=True, size=16)

    cp = doc.core_properties
    cp.title = title
    cp.author = "Qin Liu; Wangyang Ye; Ming Luo"
    cp.subject = "Mendelian randomization; falsification framework"
    cp.keywords = ("Alzheimer's disease; APOE; benchmarking; causal inference; "
                   "falsification framework; locus-driven confounding; "
                   "Mendelian randomization; operating characteristics")

    for ln in lines[1:]:
        s = ln.strip()
        if not s:
            continue
        if s.lower().startswith("## abstract"):
            break
        if s.startswith("**Qin Liu") or "Ming Luo" in s and s.startswith("**"):
            p = para(doc, "", space_after=4, line=WD_LINE_SPACING.SINGLE)
            write_authors(p)
            continue
        if s.startswith("^1^"):
            body = s[3:]
            p = para(doc, "", space_after=4, line=WD_LINE_SPACING.SINGLE)
            r = p.add_run("1")
            r.font.superscript = True
            r.font.size = Pt(11)
            add_runs(p, " " + body, size=11, italic=True)
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            continue
        if s.startswith("†"):
            p = para(doc, "", space_after=4, line=WD_LINE_SPACING.SINGLE)
            add_runs(p, s, size=11)
            continue
        if "Correspondence" in s:
            p = para(doc, "", space_after=14, line=WD_LINE_SPACING.SINGLE)
            add_runs(p, s.replace("\\*", "*"), size=11)
            continue

    # ---- body ------------------------------------------------------------
    start = next(i for i, l in enumerate(lines)
                 if l.strip().lower().startswith("## abstract"))
    i = start
    prev_table = False
    in_refs = False
    table_no = 0

    while i < len(lines):
        raw = lines[i]
        s = raw.strip()

        if not s:
            i += 1
            continue
        if s == "---":
            i += 1
            continue

        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            table_no += 1
            add_table(doc, parse_table(block), FIXED_WIDTHS.get(table_no))
            prev_table = True
            continue

        if s.startswith("#"):
            level = len(s) - len(s.lstrip("#"))
            heading = s[level:].strip()
            in_refs = heading.lower().startswith("references")
            if level == 2:
                p = para(doc, "", space_before=18, space_after=8,
                         line=WD_LINE_SPACING.SINGLE)
                add_runs(p, heading, bold=True, size=14)
            elif level == 3:
                p = para(doc, "", space_before=14, space_after=6,
                         line=WD_LINE_SPACING.SINGLE)
                add_runs(p, heading, bold=True, italic=True, size=12)
            else:
                p = para(doc, "", space_before=12, space_after=6,
                         line=WD_LINE_SPACING.SINGLE)
                add_runs(p, heading, bold=True, size=13)
            prev_table = False
            i += 1
            continue

        # table notes directly after a table
        if prev_table and s.startswith("*") and not s.startswith("**"):
            p = para(doc, "", space_before=0, space_after=14,
                     line=WD_LINE_SPACING.SINGLE)
            add_runs(p, s, size=NOTE_SIZE)
            prev_table = False
            i += 1
            continue

        # numbered protocol steps
        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m and not in_refs:
            p = para(doc, "", space_after=6, indent_cm=0.75, hanging_cm=0.75)
            add_runs(p, m.group(1) + ".  ")
            add_runs(p, m.group(2))
            prev_table = False
            i += 1
            continue

        # bullet lists
        if s.startswith("- "):
            p = para(doc, "", space_after=4, indent_cm=0.75, hanging_cm=0.5)
            add_runs(p, "•  ")
            add_runs(p, s[2:])
            prev_table = False
            i += 1
            continue

        # references: hanging indent, no extra space
        if in_refs:
            p = para(doc, "", space_after=6, indent_cm=0.8, hanging_cm=0.8)
            add_runs(p, s)
            i += 1
            continue

        # table / figure captions
        if s.startswith("**Table ") or s.startswith("**Figure "):
            p = para(doc, "", space_before=14, space_after=6,
                     line=WD_LINE_SPACING.SINGLE, keep_with_next=True)
            add_runs(p, s)
            prev_table = False
            i += 1
            continue

        p = para(doc, "", space_after=6)
        add_runs(p, s)
        prev_table = False
        i += 1

    doc.save(out_path)
    return out_path


def write_authors(p):
    """Author line with pandoc-style superscripts, e.g. Name^1,†^."""
    text = ("Qin Liu", "1,†"), ("Wangyang Ye", "1,†"), ("Ming Luo", "1,*")
    for k, (name, sup) in enumerate(text):
        r = p.add_run(name)
        r.bold = True
        r.font.name = BODY_FONT
        _east_asia(r, BODY_FONT)
        r.font.size = Pt(12)
        rs = p.add_run(sup)
        rs.bold = True
        rs.font.superscript = True
        rs.font.name = BODY_FONT
        _east_asia(rs, BODY_FONT)
        rs.font.size = Pt(12)
        if k < len(text) - 1:
            r2 = p.add_run(", ")
            r2.bold = True
            r2.font.name = BODY_FONT
            _east_asia(r2, BODY_FONT)
            r2.font.size = Pt(12)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="inp", required=True)
    ap.add_argument("--out", dest="out", required=True)
    args = ap.parse_args()
    if not os.path.exists(args.inp):
        sys.exit("input not found: " + args.inp)
    path = build(args.inp, args.out)
    print("wrote", path)


if __name__ == "__main__":
    main()
