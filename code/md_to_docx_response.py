#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build a submission-ready DOCX from the point-by-point response letter.

Layout differs from the manuscript file on purpose: reviewers read this on
screen and scroll, so the letter uses 1.5 line spacing, no line numbers, and a
visually distinct treatment for the three block types that matter — the quoted
comment, the authors' response, and the list of changes.

* The reviewer's comment is quoted in italics with a left rule, so that the
  editor can see at a glance that every point was answered in its own words.
* Response paragraphs are set normally; the "Response:" verdict stays bold.
* The "Changes:" line is smaller and grey, because it is an index, not an
  argument.
* Tables are three-line tables without vertical rules, as R3-12 asks.

Pure python-docx, no pandoc, no LibreOffice.

Usage
    python code/md_to_docx_response.py \
        --in  reviews/Response_to_Reviewers_Frontiers.md \
        --out reviews/Response_to_Reviewers_Frontiers.docx
"""

import argparse
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
BODY_SIZE = 11.5
QUOTE_SIZE = 10.5
CHANGES_SIZE = 9.5
TABLE_SIZE = 9
GREY = RGBColor(0x44, 0x44, 0x44)
TEXT_WIDTH_CM = 15.9

# Fixed widths for the disposition table: the comment id and the verdict are
# short and must not be inflated by the length of the topic text.
FIXED_WIDTHS = {4: [1.55, 5.45, 3.10, 5.80]}


# --------------------------------------------------------------------------
# low-level helpers
# --------------------------------------------------------------------------

def _set_run_font(run, name=BODY_FONT, size=None, bold=None, italic=None,
                  colour=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if colour is not None:
        run.font.color.rgb = colour


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


def add_runs(p, text, size=BODY_SIZE, bold=False, italic=False, colour=None):
    """Write inline markdown as Word runs: **bold**, *italic*, `code`."""
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = p.add_run(tok[2:-2])
            _set_run_font(r, size=size, bold=True, italic=italic, colour=colour)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = p.add_run(tok[1:-1])
            _set_run_font(r, name=MONO_FONT, size=size - 1,
                          bold=bold, italic=False, colour=colour)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = p.add_run(tok[1:-1])
            _set_run_font(r, size=size, bold=bold, italic=True, colour=colour)
        else:
            r = p.add_run(tok)
            _set_run_font(r, size=size, bold=bold, italic=italic, colour=colour)
    return p


def para(doc, space_before=0, space_after=6, line=1.5, align=None,
         left_cm=None, hanging_cm=None, keep_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if line == WD_LINE_SPACING.SINGLE:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    else:
        pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if left_cm is not None:
        pf.left_indent = Cm(left_cm)
    if hanging_cm is not None:
        pf.first_line_indent = Cm(-hanging_cm)
    if keep_next:
        pf.keep_with_next = True
    return p


def left_rule(p, sz=18, colour="999999"):
    """A vertical rule at the left edge, the usual way to mark a quotation."""
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(sz))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), colour)
    bdr.append(left)
    pPr.append(bdr)


def page_footer(section, text="Page "):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run(text), size=9, colour=GREY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


# --------------------------------------------------------------------------
# tables
# --------------------------------------------------------------------------

def _borders_none(table):
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(el)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.insert_element_before(borders, "w:shd", "w:tblLayout",
                                "w:tblCellMar", "w:tblLook")


def _cell_border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = OxmlElement("w:" + edge)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000")
    borders.append(el)


def parse_table(lines):
    rows = []
    for ln in lines:
        s = ln.strip()
        if not s.startswith("|"):
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue
        rows.append(cells)
    return rows


def column_widths(rows, total_cm=TEXT_WIDTH_CM, min_cm=1.6):
    ncols = max(len(r) for r in rows)
    weights = []
    for c in range(ncols):
        lens = [len(re.sub(r"[*`]", "", r[c])) if c < len(r) else 0 for r in rows]
        weights.append(max(2.0, sum(lens) / max(1, len(lens))) ** 0.75)
    total = sum(weights)
    widths = [max(min_cm, total_cm * w / total) for w in weights]
    scale = total_cm / sum(widths)
    return [w * scale for w in widths]


def add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _borders_none(table)

    widths = FIXED_WIDTHS.get(ncols) or column_widths(rows)

    tblPr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(round(sum(widths) * 567))))
    tblPr.insert_element_before(tblW, "w:jc", "w:tblCellSpacing", "w:tblInd",
                                "w:tblBorders", "w:shd", "w:tblLayout",
                                "w:tblCellMar", "w:tblLook")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.insert_element_before(layout, "w:tblCellMar", "w:tblLook")

    for ri, rowdata in enumerate(rows):
        row = table.rows[ri]
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        cant.set(qn("w:val"), "true")
        trPr.append(cant)
        if ri == 0:
            hdr = OxmlElement("w:tblHeader")
            hdr.set(qn("w:val"), "true")
            trPr.append(hdr)
        for ci, val in enumerate(rowdata):
            cell = row.cells[ci]
            cell.width = Cm(widths[ci])
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.space_before = Pt(2)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.0
            add_runs(p, val, size=TABLE_SIZE, bold=(ri == 0))

    for cell in table.rows[0].cells:
        _cell_border(cell, "top", 12)
        _cell_border(cell, "bottom", 8)
    for cell in table.rows[-1].cells:
        _cell_border(cell, "bottom", 12)

    after = para(doc, space_before=0, space_after=10,
                 line=WD_LINE_SPACING.SINGLE)
    after.paragraph_format.line_spacing = Pt(6)
    return table


# --------------------------------------------------------------------------
# conversion
# --------------------------------------------------------------------------

def setup(doc):
    for s in doc.sections:
        s.page_height = Cm(29.7)
        s.page_width = Cm(21.0)
        s.top_margin = Cm(2.54)
        s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(2.54)
        s.right_margin = Cm(2.54)
        page_footer(s)
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(BODY_SIZE)


def build(in_path, out_path):
    lines = open(in_path, encoding="utf-8").read().split("\n")
    doc = Document()
    setup(doc)

    cp = doc.core_properties
    cp.title = "Response to Reviewers — manuscript 1986036"
    cp.author = "Qin Liu; Wangyang Ye; Ming Luo"
    cp.subject = "Frontiers in Genetics, manuscript 1986036"
    cp.keywords = ("response to reviewers; Mendelian randomization; "
                   "falsification framework")

    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i]
        s = raw.strip()
        i += 1
        if not s or s == "---":
            continue

        # headings
        if s.startswith("### "):
            p = para(doc, space_before=12, space_after=4, line=1.15,
                     keep_next=True)
            add_runs(p, s[4:], size=11.5, bold=True)
            continue
        if s.startswith("## "):
            p = para(doc, space_before=16, space_after=5, line=1.15,
                     keep_next=True)
            add_runs(p, s[3:], size=13, bold=True)
            continue
        if s.startswith("# "):
            p = para(doc, space_before=0, space_after=12, line=1.15)
            add_runs(p, s[2:], size=16, bold=True)
            continue

        # quoted reviewer comment: contiguous '>' lines, may wrap
        if s.startswith(">"):
            block = [s.lstrip("> ").rstrip()]
            while i < n and lines[i].strip().startswith(">"):
                block.append(lines[i].strip().lstrip("> ").rstrip())
                i += 1
            text = " ".join(x for x in block if x)
            text = text.strip()
            if text.startswith("*") and text.endswith("*"):
                text = text[1:-1]
            p = para(doc, space_before=6, space_after=8, line=1.25,
                     left_cm=0.7)
            p.paragraph_format.right_indent = Cm(0.4)
            add_runs(p, text, size=QUOTE_SIZE, italic=True, colour=GREY)
            left_rule(p)
            continue

        # tables
        if s.startswith("|"):
            block = [s]
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            add_table(doc, parse_table(block))
            continue

        # the changes index line
        if s.startswith("**Changes:**"):
            p = para(doc, space_before=4, space_after=10, line=1.2)
            add_runs(p, s, size=CHANGES_SIZE, colour=GREY)
            continue

        # numbered list
        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            p = para(doc, space_before=0, space_after=6, line=1.35,
                     align=WD_ALIGN_PARAGRAPH.JUSTIFY, left_cm=0.75,
                     hanging_cm=0.75)
            add_runs(p, s, size=BODY_SIZE)
            continue

        # body paragraph
        p = para(doc, space_before=0, space_after=7, line=1.5,
                 align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        add_runs(p, s, size=BODY_SIZE)

    doc.save(out_path)
    words = sum(len(p.text.split()) for p in doc.paragraphs)
    print(f"wrote {out_path}")
    print(f"  paragraphs {len(doc.paragraphs)}  tables {len(doc.tables)}  words {words:,}")
    return 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="src", required=True)
    ap.add_argument("--out", dest="dst", required=True)
    a = ap.parse_args()
    if not os.path.exists(a.src):
        raise SystemExit(f"missing input: {a.src}")
    return build(a.src, a.dst)


if __name__ == "__main__":
    raise SystemExit(main())
